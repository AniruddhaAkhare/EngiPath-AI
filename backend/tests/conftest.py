"""
EngiPath AI — Test Configuration

Shared fixtures for all test modules.
"""
import os
import io
import pytest

# Set testing env vars before importing app
os.environ.update({
    "FLASK_ENV": "testing",
    "DATABASE_URL": "sqlite:///:memory:",
    "CLASSES_GEMINI_API_KEY": "test-classes-key-abc123",
    "CLASSES_GEMINI_MODEL": "gemini-2.5-flash",
    "INTERNSHIP_GEMINI_API_KEY": "test-internship-key-xyz789",
    "INTERNSHIP_GEMINI_MODEL": "gemini-flash-latest",
    "SEARCH_API_KEY": "test-search-key",
    "SEARCH_ENGINE_ID": "test-engine-id",
    "UPLOAD_FOLDER": "uploads_test",
    "DELETE_RESUME_AFTER_ANALYSIS": "false",
    "SCRAPER_REQUEST_TIMEOUT": "5",
})

from app import create_app
from app.extensions import db as _db


@pytest.fixture(scope="session")
def app():
    """Create application for the test session."""
    import os
    os.makedirs("uploads_test", exist_ok=True)
    application = create_app("testing")
    with application.app_context():
        _db.create_all()
        yield application
        _db.session.remove()
        _db.drop_all()


@pytest.fixture(scope="function")
def client(app):
    """Create a test client for each test function."""
    with app.test_client() as c:
        yield c


@pytest.fixture(scope="function")
def db(app):
    """Provide database access for each test with rollback cleanup."""
    yield _db
    _db.session.rollback()


@pytest.fixture
def mock_pdf_file():
    """A minimal fake PDF-like byte stream for upload testing."""
    # Minimal PDF header
    content = b"%PDF-1.4\n1 0 obj\n<< /Type /Catalog >>\nendobj\n%%EOF"
    return (io.BytesIO(content), "test_resume.pdf", "application/pdf")


@pytest.fixture
def mock_docx_file():
    """Creates an actual minimal DOCX file in memory."""
    from docx import Document
    doc = Document()
    doc.add_paragraph("John Doe")
    doc.add_paragraph("Skills: Python, Machine Learning, SQL")
    doc.add_paragraph("Education: B.Tech Computer Science, 2024")
    doc.add_paragraph("Projects: Sentiment Analysis using BERT")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return (buf, "test_resume.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@pytest.fixture
def sample_class_data():
    """Sample class data for comparison tests."""
    return {
        "institute_name": "Test Institute",
        "course_name": "Python Programming",
        "address": "123 Test Street",
        "city": "Pune",
        "state": "Maharashtra",
        "country": "India",
        "fees": 15000.0,
        "currency": "INR",
        "duration": 3.0,
        "duration_unit": "months",
        "mode": "offline",
        "description": "Comprehensive Python course",
        "topics": ["Python", "OOP", "Django"],
        "skills": ["Python", "Web Development"],
        "website_url": "https://testinstitute.example.com",
        "source_url": "https://example.com/source",
    }
